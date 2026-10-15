"""Shared visual renderers for PDF, DOCX and PPTX exports."""
import io
import re
import textwrap
from xml.sax.saxutils import escape
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import markdown as mdlib
from bs4 import BeautifulSoup
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage,
    PageBreak, KeepTogether
)
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

DEFAULT_PALETTE = ["#2563EB", "#14B8A6", "#F59E0B", "#EF4444", "#8B5CF6", "#06B6D4"]

def markdown_to_html(markdown_text: str) -> str:
    return mdlib.markdown(markdown_text, extensions=["tables", "fenced_code", "sane_lists"])

def _hex_color(value, fallback="#2563EB"):
    if not isinstance(value, str) or not re.fullmatch(r"#[0-9a-fA-F]{6}", value):
        value = fallback
    return value

def _rgb(hex_color):
    h = _hex_color(hex_color).lstrip("#")
    return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))

def _theme_from_query(query: str = "") -> dict:
    q = (query or "").lower()
    choices = {
        "blue": ["#1D4ED8", "#60A5FA", "#0F172A", "#DBEAFE"],
        "green": ["#15803D", "#4ADE80", "#14532D", "#DCFCE7"],
        "purple": ["#7C3AED", "#C4B5FD", "#4C1D95", "#EDE9FE"],
        "orange": ["#EA580C", "#FDBA74", "#7C2D12", "#FFEDD5"],
        "red": ["#DC2626", "#FCA5A5", "#7F1D1D", "#FEE2E2"],
        "teal": ["#0F766E", "#5EEAD4", "#134E4A", "#CCFBF1"],
        "dark": ["#0F172A", "#64748B", "#F8FAFC", "#E2E8F0"],
    }
    for name, palette in choices.items():
        if name in q and any(word in q for word in ("color", "colour", "theme", "palette", "attractive", "style")):
            return {"primary": palette[0], "secondary": palette[1],
                    "text": palette[2], "light": palette[3], "palette": palette}
    return {"primary": DEFAULT_PALETTE[0], "secondary": DEFAULT_PALETTE[1],
            "text": "#0F172A", "light": "#DBEAFE", "palette": DEFAULT_PALETTE}

def _parse_markdown_blocks(markdown_text: str):
    """Parse headings, paragraphs, and pipe tables in their original order."""
    lines = markdown_text.splitlines()
    blocks, paragraph, i = [], [], 0
    def flush():
        nonlocal paragraph
        if paragraph:
            blocks.append(("text", "\n".join(paragraph).strip()))
            paragraph = []
    while i < len(lines):
        line = lines[i]
        if "|" in line and i + 1 < len(lines) and re.search(r"\|?\s*:?-{3,}", lines[i+1]):
            flush()
            headers = [x.strip() for x in line.strip().strip("|").split("|")]
            rows, i = [], i + 2
            while i < len(lines) and "|" in lines[i]:
                cells = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if len(cells) == len(headers):
                    rows.append(cells)
                i += 1
            if headers:
                blocks.append(("table", [headers] + rows))
            continue
        if not line.strip():
            flush()
        else:
            paragraph.append(line)
        i += 1
    flush()
    return blocks

def _plain(s):
    s = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", s)
    s = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", s)
    s = re.sub(r"(\*\*|__|`|~~|\*)", "", s)
    return s.strip()

def _chart_png(spec: dict, dpi=160, theme=None) -> bytes:
    theme = theme or _theme_from_query()
    kind, title = spec["type"], spec.get("title", "Chart")
    fig, ax = plt.subplots(figsize=(9, 5.0), constrained_layout=True)
    palette = [_hex_color(c) for c in (spec.get("colors") or theme["palette"])]
    labels, values = spec.get("labels", []), spec.get("values", [])
    if kind == "pie":
        wedges, texts, autotexts = ax.pie(
            values, labels=labels, autopct=lambda p: f"{p:.1f}%" if p >= 3 else "",
            startangle=90, colors=[palette[i % len(palette)] for i in range(len(values))],
            wedgeprops={"edgecolor":"white", "linewidth":1.5}
        )
        ax.axis("equal")
    elif kind == "histogram":
        ax.hist(values, bins=min(12, max(3, int(np.sqrt(len(values))))),
                color=palette[0], edgecolor="white", alpha=0.9)
        ax.set_xlabel(spec.get("x_label") or "Value")
        ax.set_ylabel("Frequency")
        ax.grid(axis="y", alpha=0.25)
    elif kind == "scatter":
        ax.scatter(spec["x_values"], spec["y_values"], c=palette[0], alpha=0.8,
                   s=65, edgecolors="white", linewidths=0.6)
        ax.set_xlabel(spec.get("x_label") or "X")
        ax.set_ylabel(spec.get("y_label") or "Y")
        ax.grid(alpha=0.25)
    elif kind in ("line", "timeline"):
        x = list(range(len(values)))
        ax.plot(x, values, color=palette[0], marker="o", linewidth=2.5)
        ax.fill_between(x, values, alpha=0.12, color=palette[0])
        ax.set_xticks(x)
        ax.set_xticklabels(labels, rotation=30, ha="right")
        ax.set_ylabel(spec.get("y_label") or "")
        ax.set_xlabel(spec.get("x_label") or ("Time" if kind == "timeline" else ""))
        ax.grid(axis="y", alpha=0.25)
    else:
        bars = ax.bar(labels, values, color=[palette[i % len(palette)] for i in range(len(values))])
        ax.bar_label(bars, fmt="%.3g", padding=3, fontsize=8)
        ax.set_xlabel(spec.get("x_label") or "")
        ax.set_ylabel(spec.get("y_label") or "")
        ax.tick_params(axis="x", labelrotation=25)
        ax.grid(axis="y", alpha=0.2)
        ax.set_axisbelow(True)
    ax.set_title(title, fontsize=15, fontweight="bold", pad=14)
    fig.patch.set_facecolor("white")
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return buf.getvalue()

def _table_to_reportlab(rows, theme):
    cooked = []
    for r, row in enumerate(rows):
        cooked.append([Paragraph(escape(_plain(str(c))), ParagraphStyle(
            f"cell{r}", fontName="Helvetica-Bold" if r == 0 else "Helvetica",
            fontSize=7.5, leading=9, textColor=colors.white if r == 0 else colors.HexColor("#111827")
        )) for c in row])
    t = Table(cooked, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor(theme["primary"])),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.35, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING", (0,0), (-1,-1), 5), ("RIGHTPADDING", (0,0), (-1,-1), 5),
        ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5),
    ]))
    return t

def _story_from_markdown(md, styles, theme, include_page_breaks=False):
    story = []
    blocks = _parse_markdown_blocks(md)
    for kind, content in blocks:
        if kind == "table":
            story.append(_table_to_reportlab(content, theme))
            story.append(Spacer(1, 8))
            continue
        raw = content.strip()
        if not raw:
            continue
        if raw.startswith("# "):
            story.append(Paragraph(escape(_plain(raw[2:])), styles["Title"]))
        elif raw.startswith("## "):
            story.append(Spacer(1, 5))
            story.append(Paragraph(escape(_plain(raw[3:])), styles["Heading2"]))
        elif raw.startswith("### "):
            story.append(Paragraph(escape(_plain(raw[4:])), styles["Heading3"]))
        else:
            text = escape(_plain(raw)).replace("\n", "<br/>")
            story.append(Paragraph(text, styles["BodyText"]))
            story.append(Spacer(1, 6))
    return story

def _build_pdf_story(md, charts, theme, dpi=160):
    styles = getSampleStyleSheet()
    styles["Title"].textColor = colors.HexColor(theme["primary"])
    styles["Heading2"].textColor = colors.HexColor(theme["primary"])
    styles["Heading3"].textColor = colors.HexColor(theme["secondary"])
    styles["BodyText"].fontSize = 9
    styles["BodyText"].leading = 12
    story = _story_from_markdown(md, styles, theme)
    # Each requested chart is explicitly placed on its own page after report content.
    for chart in charts:
        story.append(PageBreak())
        story.append(Paragraph(escape(chart.get("title", "Chart")), styles["Heading1"]))
        story.append(Spacer(1, 8))
        story.append(RLImage(io.BytesIO(_chart_png(chart, dpi, theme)), width=175*mm, height=98*mm, kind="proportional"))
        if chart.get("notes"):
            story.append(Spacer(1, 5))
            story.append(Paragraph(escape(chart["notes"]), styles["BodyText"]))
    return story

def export_pdf(markdown_text: str, charts=None, query: str = "") -> bytes:
    charts = charts or []
    theme = _theme_from_query(query)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=16*mm, leftMargin=16*mm,
                            topMargin=16*mm, bottomMargin=16*mm, title="AI Generated Report")
    story = _build_pdf_story(markdown_text, charts, theme)
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(16*mm, 9*mm, "AI Report Generator")
        canvas.drawRightString(A4[0]-16*mm, 9*mm, f"Page {doc.page}")
        canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()

def export_docx(markdown_text: str, charts=None, query: str = "") -> bytes:
    charts = charts or []
    theme = _theme_from_query(query)
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Aptos", Pt(10)
    normal.font.color.rgb = RGBColor(*_rgb(theme["text"]))
    for kind, content in _parse_markdown_blocks(markdown_text):
        if kind == "table":
            table = doc.add_table(rows=1, cols=len(content[0]))
            table.style = "Light Shading Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, value in enumerate(content[0]):
                table.rows[0].cells[i].text = _plain(value)
            for row in content[1:]:
                cells = table.add_row().cells
                for i, value in enumerate(row):
                    cells[i].text = _plain(value)
            doc.add_paragraph()
        else:
            text = content.strip()
            if text.startswith("# "):
                doc.add_heading(_plain(text[2:]), 0)
            elif text.startswith("## "):
                doc.add_heading(_plain(text[3:]), 1)
            elif text.startswith("### "):
                doc.add_heading(_plain(text[4:]), 2)
            else:
                doc.add_paragraph(_plain(text))
    for chart in charts:
        doc.add_page_break()
        doc.add_heading(chart.get("title", "Chart"), level=1)
        doc.add_picture(io.BytesIO(_chart_png(chart, theme=theme)), width=Inches(6.1))
        if chart.get("notes"):
            doc.add_paragraph(chart["notes"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def _set_cell_fill(cell, hex_color):
    fill = cell.fill
    fill.solid()
    r, g, b = _rgb(hex_color)
    fill.fore_color.rgb = PptRGB(r, g, b)

def _ppt_add_text(slide, text, x, y, w, h, size=16, bold=False, color="#0F172A"):
    shape = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.text = text
    for p in tf.paragraphs:
        for run in p.runs:
            run.font.size = PptPt(size)
            run.font.bold = bold
            run.font.color.rgb = PptRGB(*_rgb(color))
    return shape

def export_pptx(markdown_text: str, charts=None, query: str = "") -> bytes:
    charts = charts or []
    theme = _theme_from_query(query)
    prs = Presentation()
    prs.slide_width, prs.slide_height = PptInches(13.333), PptInches(7.5)
    blank = prs.slide_layouts[6]
    # Main report slide(s): actual PowerPoint tables, not pipe-delimited text.
    slide = prs.slides.add_slide(blank)
    _ppt_add_text(slide, "Report", 0.45, 0.22, 12.2, 0.5, 26, True, theme["primary"])
    y = 0.9
    for kind, content in _parse_markdown_blocks(markdown_text):
        if kind == "table":
            rows, cols = len(content), len(content[0])
            # Start a new slide if a table would exceed the current slide.
            estimated_h = min(3.8, 0.28 * (rows + 1))
            if y + estimated_h > 6.85:
                slide = prs.slides.add_slide(blank)
                y = 0.45
            shape = slide.shapes.add_table(rows, cols, PptInches(0.45), PptInches(y),
                                           PptInches(12.4), PptInches(estimated_h))
            table = shape.table
            for r, row in enumerate(content):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = _plain(str(value))[:500]
                    cell.vertical_anchor = 1
                    if r == 0:
                        _set_cell_fill(cell, theme["primary"])
                    else:
                        _set_cell_fill(cell, "#F1F5F9" if r % 2 else "#FFFFFF")
                    for p in cell.text_frame.paragraphs:
                        for run in p.runs:
                            run.font.size = PptPt(10 if r == 0 else 9)
                            run.font.bold = r == 0
                            run.font.color.rgb = PptRGB(255,255,255) if r == 0 else PptRGB(15,23,42)
            y += estimated_h + 0.2
        else:
            text = content.strip()
            if not text:
                continue
            if text.startswith("# "):
                heading, size = _plain(text[2:]), 22
            elif text.startswith("## "):
                heading, size = _plain(text[3:]), 16
            elif text.startswith("### "):
                heading, size = _plain(text[4:]), 13
            else:
                heading, size = _plain(text), 11
            lines = max(1, min(4, (len(heading) // 115) + 1))
            h = 0.35 if size >= 16 else min(0.65, 0.22 * lines + 0.12)
            if y + h > 6.9:
                slide = prs.slides.add_slide(blank)
                y = 0.45
            _ppt_add_text(slide, heading, 0.55, y, 12.0, h, size, size >= 16,
                          theme["primary"] if size >= 16 else theme["text"])
            y += h + 0.08
    # One full chart per slide, consistent with PDF's chart-page layout.
    for chart in charts:
        slide = prs.slides.add_slide(blank)
        _ppt_add_text(slide, chart.get("title", "Chart"), 0.5, 0.25, 12.3, 0.55,
                      24, True, theme["primary"])
        slide.shapes.add_picture(io.BytesIO(_chart_png(chart, theme=theme)),
                                 PptInches(1.0), PptInches(1.0),
                                 width=PptInches(11.3), height=PptInches(5.7))
        if chart.get("notes"):
            _ppt_add_text(slide, chart["notes"], 0.65, 6.85, 12.0, 0.35, 9, False, theme["text"])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
